#!/usr/bin/env python3
"""Validate a generated workbook against the locked task specification."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--spec", required=True)
    parser.add_argument("--report", required=True)
    args = parser.parse_args()
    xlsx = Path(args.xlsx).resolve()
    spec_path = Path(args.spec).resolve()
    report_path = Path(args.report).resolve()
    errors: list[dict] = []
    warnings: list[dict] = []
    info: list[dict] = []

    try:
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
        wb = load_workbook(xlsx, read_only=True, data_only=False)
        visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
        if not visible:
            errors.append({"code": "NO_VISIBLE_WORKSHEET"})
        text_parts: list[str] = []
        nonempty = 0
        formulas = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    nonempty += 1
                    text_parts.append(str(value))
                    if isinstance(value, str) and value.startswith("="):
                        formulas += 1
        wb.close()
        if nonempty == 0:
            errors.append({"code": "EMPTY_WORKBOOK"})
        searchable = "\n".join(text_parts).casefold()
        for requirement in spec.get("quote_requirements", []):
            if not isinstance(requirement, dict) or not requirement.get("required"):
                continue
            if requirement.get("status") not in {"confirmed", "derived"}:
                continue
            keywords = [str(value).strip() for value in requirement.get("keywords", []) if str(value).strip()]
            if not keywords:
                warnings.append({"code": "QUOTE_REQUIREMENT_WITHOUT_KEYWORDS",
                                 "requirement_id": requirement.get("id")})
                continue
            if not any(keyword.casefold() in searchable for keyword in keywords):
                errors.append({"code": "MISSING_REQUIRED_QUOTE_ITEM",
                               "requirement_id": requirement.get("id"), "keywords": keywords})
        info.append({"code": "WORKBOOK_STATS", "sheets": wb.sheetnames,
                     "nonempty_cells": nonempty, "formula_cells": formulas})
    except Exception as exc:
        errors.append({"code": "XLSX_VALIDATOR_ERROR", "exception_type": type(exc).__name__,
                       "message": str(exc)})

    report = {"schema": 1, "status": "failed" if errors else "passed",
              "xlsx": str(xlsx), "model_spec": str(spec_path),
              "errors": errors, "warnings": warnings, "info": info}
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"XLSX_VALIDATION_REPORT={report_path}")
    print(f"XLSX_VALIDATION_STATUS={report['status']} errors={len(errors)} warnings={len(warnings)}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
